from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin    = Pt(60)
    section.bottom_margin = Pt(60)
    section.left_margin   = Pt(72)
    section.right_margin  = Pt(72)

def para(text='', bold=False, italic=False, size=11, align=None, color=None, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = 'Arial'
        if color:
            run.font.color.rgb = RGBColor(*color)
    return p

def highlight_heading(text, hex_color='F4B942'):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    rPr.append(shd)
    return p

def sys_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    return p

def bullet(text, bold_part='', normal_part=''):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(3)
    if bold_part:
        r1 = p.add_run(bold_part)
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.name = 'Arial'
        if normal_part:
            r2 = p.add_run(normal_part)
            r2.font.size = Pt(11)
            r2.font.name = 'Arial'
    else:
        r = p.add_run(text)
        r.font.size = Pt(11)
        r.font.name = 'Arial'
    return p

# ── Заголовок ──
para('Заключение по тестированию', bold=True, size=14,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('06.10.2026', size=11, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=6)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run('Josef Urmann, 12.03.1956')
r.bold = True; r.font.size = Pt(11); r.font.name = 'Arial'
para('Биологический возраст — 75 лет', italic=True, size=11, space_after=12)

# ── Заголовок раздела ──
highlight_heading('Что происходит в организме:')

# ── Голова и нервы ──
sys_title('Голова и нервы')
bullet('Мозг справа немного перегружен — из-за этого бывает сложно сдерживать эмоции и реакции')
bullet('Нервная система быстро устаёт и долго не может успокоиться')

# ── Позвоночник и суставы ──
sys_title('Позвоночник и суставы')
bullet('Два верхних позвонка в шее сдвинуты — они пережимают сосуды и нервы. Из-за этого давление скачет, немеют руки, кружится голова, хуже слышно и видно, нарушается координация')
bullet('В середине и внизу спины диски немного выпирают')
bullet('Ноги ощущаются слабыми')
bullet('Плечо воспалено — болит сустав и связки')
bullet('Маленькие суставы на ногах воспалены')
bullet('Мышца между грудью и животом стоит не совсем на своём месте')
bullet('Плоскостопие')
bullet('Кости стали менее плотными — легче ломаются')

# ── Сердце и сосуды ──
sys_title('Сердце и сосуды')
bullet('На стенках сосудов осели жировые отложения — сосуды стали уже')
bullet('Когда есть стресс — в крови резко повышается жир')
bullet('Вены на ногах начинают расширяться')

# ── Живот и пищеварение ──
sys_title('Живот и пищеварение')
bullet('Желчный пузырь воспалён, желчь плохо вытекает — могут образовываться камни')
bullet('Поджелудочная железа раздражена и работает не в полную силу')
bullet('В животе часто скапливаются газы')

# ── Гормоны ──
sys_title('Гормоны')
bullet('Щитовидная железа работает вяло')
bullet('Железы, которые управляют кальцием в организме, дают сбои')
bullet('Половые гормоны снижены')
bullet('Иммунитет иногда путается и начинает атаковать свои же клетки')
bullet('Бывают резкие волны адреналина — сердце колотится, появляется тревога, давление скачет')

# ── Мочевой пузырь и простата ──
sys_title('Мочевой пузырь и простата')
bullet('Мочевой пузырь плохо сокращается — частые позывы или ощущение, что не до конца опустошился')
bullet('Простата увеличена')

# ── Глаза ──
sys_title('Глаза')
bullet('В левом глазу изменились сосуды на сетчатке')
bullet('В правом глазу хрусталик начинает мутнеть, центральное зрение снижается')

# ── Нос, горло, зубы ──
sys_title('Нос, горло, зубы')
bullet('Носовая перегородка кривая — трудно дышать носом')
bullet('Носовые пазухи хронически воспалены')
bullet('Миндалины хронически воспалены')
bullet('Мягкое нёбо расслаблено — из-за этого человек храпит и плохо спит')
bullet('Дёсны воспалены')

# ── Инфекции ──
sys_title('Инфекции')
bullet('Грибок (молочница) в мочеиспускательном канале, в ушах снаружи и внутри')
bullet('В организме активен вирус герпеса 1-го, 2-го и 7-го типа')
bullet('Недавно (в последние 3 месяца) перенёс COVID')
bullet('В организме есть паразиты')

# ── Чего не хватает ──
sys_title('Чего не хватает организму')
bullet('Не хватает витаминов А, D и группы В')
bullet('Не хватает минералов и цинка')
bullet('Не хватает строительного материала для клеток — аминокислот: аргинин, глицин, аспарагиновая кислота, таурин, орнитин, диметилглицин, глютатион, валин, метионин')
bullet('В тканях накопились токсины от плесени')

# ── Что плохо переносит ──
sys_title('Что плохо переносит организм')
bullet('Плохо реагирует на: бобовые, сладкое, рис, всё из пшеницы / ржи / овса, жирную рыбу')
bullet('Чувствительность к запахам: духи, выхлопные газы и похожие вещества')

# ── Подвал ──
doc.add_paragraph()
for line in [
    'Поздравляю!',
    'Благодарю за доверие.',
    '',
    'Я рада приветствовать Вас на Вашей программе коррекции веса и здоровья.',
    '',
    'Моя цель — помочь Вам справиться с Вашей проблемой и научиться осознанно управлять своим здоровьем.',
    '',
    'С заботой и верой в Вас, ваш Health-coach Anna Soloviova и команда проекта WELLNESS CODE by Anna Soloviova',
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(line)
    run.italic = True
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

doc.save('/home/user/Soloviova/Urmann_Josef_simple.docx')
print('Done')
