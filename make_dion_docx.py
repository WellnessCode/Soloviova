"""Generate Dion Vilhelmsen patient summary in Wellness Code format."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PURPLE_HEX   = '7B52AB'
LAVENDER_HEX = 'EEE8F8'
PURPLE_RGB   = RGBColor(123, 82, 171)
ORANGE_RGB   = RGBColor(240, 129, 30)
BLACK_RGB    = RGBColor(30, 30, 30)
FONT = 'Calibri'


def shd(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'),   'clear')
    e.set(qn('w:color'), 'auto')
    e.set(qn('w:fill'),  hex_color)
    tcPr.append(e)


def get_or_add_tblPr(tbl_elem):
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl_elem.insert(0, tblPr)
    return tblPr


def no_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl)
    tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'), 'none')
        tb.append(b)
    tblPr.append(tb)


def cell_margins(cell, top=0, bottom=0, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in [('top', top),('bottom', bottom),('left', left),('right', right)]:
        e = OxmlElement(f'w:{name}')
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcPr.append(mar)


def spacing(para, before=0, after=0, line=None):
    pPr = para._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'),  str(after))
    if line:
        sp.set(qn('w:line'), str(line))
        sp.set(qn('w:lineRule'), 'auto')
    pPr.append(sp)


def add_run(para, text, bold=False, color=None, size=11, italic=False):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color or BLACK_RGB
    return run


def section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    no_borders(tbl)
    tbl.rows[0].height = Cm(0.9)
    cell = tbl.cell(0, 0)
    shd(cell, PURPLE_HEX)
    cell_margins(cell, top=80, bottom=80, left=160, right=160)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacing(p, before=0, after=0)
    add_run(p, text, bold=True, color=RGBColor(255,255,255), size=12)
    doc.add_paragraph()


def info_row(doc, label, value):
    tbl = doc.add_table(rows=1, cols=2)
    no_borders(tbl)
    tbl.columns[0].width = Cm(6)
    tbl.columns[1].width = Cm(10)
    cell_margins(tbl.cell(0,0), top=40, bottom=40, left=160, right=80)
    cell_margins(tbl.cell(0,1), top=40, bottom=40, left=80, right=160)
    p0 = tbl.cell(0,0).paragraphs[0]
    spacing(p0, before=0, after=0)
    add_run(p0, label, bold=True, color=PURPLE_RGB, size=10.5)
    p1 = tbl.cell(0,1).paragraphs[0]
    spacing(p1, before=0, after=0)
    add_run(p1, value, size=10.5)


def highlight_row(doc, label, value):
    tbl = doc.add_table(rows=1, cols=2)
    no_borders(tbl)
    tbl.columns[0].width = Cm(6)
    tbl.columns[1].width = Cm(10)
    shd(tbl.cell(0,0), LAVENDER_HEX)
    shd(tbl.cell(0,1), LAVENDER_HEX)
    cell_margins(tbl.cell(0,0), top=60, bottom=60, left=160, right=80)
    cell_margins(tbl.cell(0,1), top=60, bottom=60, left=80, right=160)
    p0 = tbl.cell(0,0).paragraphs[0]
    spacing(p0, before=0, after=0)
    add_run(p0, label, bold=True, color=PURPLE_RGB, size=10.5)
    p1 = tbl.cell(0,1).paragraphs[0]
    spacing(p1, before=0, after=0)
    add_run(p1, value, bold=True, color=ORANGE_RGB, size=10.5)


def gap(doc, space=40):
    p = doc.add_paragraph()
    spacing(p, before=0, after=space)


# ── BUILD DOCUMENT ────────────────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin   = Cm(2)
    section.right_margin  = Cm(2)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p, before=0, after=100)
add_run(p, 'ВЫПИСКА ИЗ АНКЕТЫ ПАЦИЕНТА', bold=True, color=PURPLE_RGB, size=16)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p2, before=0, after=200)
add_run(p2, 'Dion Vilhelmsen', bold=True, color=ORANGE_RGB, size=13)

# ── 1. ЛИЧНЫЕ ДАННЫЕ ─────────────────────────────────────────────────────────
section_header(doc, '1. ЛИЧНЫЕ ДАННЫЕ')
info_row(doc, 'Имя:', 'Dion Vilhelmsen')
info_row(doc, 'Дата рождения:', '05.10.1975 (50 лет)')
info_row(doc, 'Телефон:', '+45 27638030')
info_row(doc, 'Email:', 'dion.vilhelmsen@gmail.com')
info_row(doc, 'Город:', 'Греве, Дания')
info_row(doc, 'Профессия:', 'Владелец бизнеса')
gap(doc)

# ── 2. АНТРОПОМЕТРИЯ ─────────────────────────────────────────────────────────
section_header(doc, '2. АНТРОПОМЕТРИЯ')
info_row(doc, 'Рост:', '178 см')
info_row(doc, 'Вес:', '86 кг')
info_row(doc, 'Объём живота:', '~100 см')
gap(doc)

# ── 3. ЦЕЛЬ ОБРАЩЕНИЯ ────────────────────────────────────────────────────────
section_header(doc, '3. ЦЕЛЬ ОБРАЩЕНИЯ')
highlight_row(doc, 'Запрос:', 'Понять работу нервной системы и пищеварения, выявить триггеры')
info_row(doc, 'Желаемый результат:', 'Снизить вес ещё на 3–4 кг')
gap(doc)

# ── 4. ОБРАЗ ЖИЗНИ ───────────────────────────────────────────────────────────
section_header(doc, '4. ОБРАЗ ЖИЗНИ')
info_row(doc, 'Физическая активность:', 'Ежедневно — упражнения, плавание, сауна (1.5–2 часа)')
info_row(doc, 'Стресс:', 'Периодически, причины известны')
info_row(doc, 'Курение:', 'Нет')
info_row(doc, 'Алкоголь:', '~2 порции')
info_row(doc, 'Сон:', 'Засыпает ~22:30, просыпается 6:00–7:00')
highlight_row(doc, 'Нарушения сна:', 'Примерно 50% ночей — нарушения')
gap(doc)

# ── 5. ПИТАНИЕ ───────────────────────────────────────────────────────────────
section_header(doc, '5. ПИТАНИЕ')
info_row(doc, 'Завтрак:', 'Пропускает или яйца с овощами (после 11:00)')
info_row(doc, 'Обед:', 'Первый приём пищи')
info_row(doc, 'Ужин:', '18:00–19:00 — мясо/курица/ягнёнок, салаты, тёплые овощи (~250–300 г)')
info_row(doc, 'Перекусы:', 'Фрукты (банан, яблоко)')
info_row(doc, 'Вода:', '1000–1500 мл/день')
info_row(doc, 'Кофе:', '2–3 чашки в день (натуральный и растворимый)')
info_row(doc, 'Сладкое:', 'Практически не употребляет (~1 порция/нед)')
highlight_row(doc, 'Чувствительность:', 'Пшеница, хлеб, молочные продукты, сахар')
info_row(doc, 'Предпочтения:', 'Тайская кухня, суши, острое и пряное')
gap(doc)

# ── 6. ПИЩЕВАРЕНИЕ ───────────────────────────────────────────────────────────
section_header(doc, '6. ПИЩЕВАРЕНИЕ')
info_row(doc, 'Стул:', '1–2 раза в день')
highlight_row(doc, 'Характер стула:', 'Жидкий / кашицеобразный / смешанный')
highlight_row(doc, 'Жалобы:', 'Повышенное газообразование, вздутие после определённых продуктов')
info_row(doc, 'Изжога:', 'Иногда по утрам и после определённых продуктов')
info_row(doc, 'Время в туалете:', 'Значительно увеличилось по сравнению с прошлым')
gap(doc)

# ── 7. НЕРВНАЯ СИСТЕМА ────────────────────────────────────────────────────────
section_header(doc, '7. НЕРВНАЯ СИСТЕМА И СОСУДЫ')
highlight_row(doc, 'Холодные руки/ноги:', 'Да — периодически')
info_row(doc, 'Стресс:', 'Есть, причины известны')
info_row(doc, 'Нервная реакция:', 'Да')
gap(doc)

# ── 8. ОПОРНО-ДВИГАТЕЛЬНЫЙ АППАРАТ ──────────────────────────────────────────
section_header(doc, '8. ОПОРНО-ДВИГАТЕЛЬНЫЙ АППАРАТ')
info_row(doc, 'Боли:', 'Колени и поясница (спортивные)')
info_row(doc, 'Осанка:', 'Периодически сутулится')
info_row(doc, 'Травмы:', 'Голеностоп, левое колено, правое ахиллово сухожилие (>15 лет назад)')
gap(doc)

# ── 9. АНАМНЕЗ ───────────────────────────────────────────────────────────────
section_header(doc, '9. АНАМНЕЗ И ЗДОРОВЬЕ')
highlight_row(doc, 'COVID-19:', 'Да, 2020 г. — тяжёлое течение')
info_row(doc, 'Паротит:', 'В детстве')
info_row(doc, 'Зрение:', 'Проблемы с правым глазом в детстве')
info_row(doc, 'Урология:', 'Одно яичко увеличено (скопление жидкости, по заключению врача)')
info_row(doc, 'Хронических заболеваний:', 'Не выявлено')
gap(doc)

# ── 10. ДОБАВКИ ──────────────────────────────────────────────────────────────
section_header(doc, '10. ДОБАВКИ И МЕДИКАМЕНТЫ')
info_row(doc, 'Добавки:', 'Комплекс качественных витаминов')
info_row(doc, 'Медикаменты:', 'Нет')
gap(doc)

# ── 11. ПРИОРИТЕТЫ ───────────────────────────────────────────────────────────
section_header(doc, '11. ПРИОРИТЕТЫ ПАЦИЕНТА')
highlight_row(doc, '№1:', 'Нервная система — выявить триггеры')
highlight_row(doc, '№2:', 'Пищеварение — газообразование, нормализация стула')
highlight_row(doc, '№3:', 'Снижение веса на 3–4 кг')

OUT = '/home/user/Soloviova/Dion_Vilhelmsen_Выписка.docx'
doc.save(OUT)
print(f'Saved: {OUT}')
