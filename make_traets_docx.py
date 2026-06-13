"""Generate Traets Magdalena patient extract in Smolych/Wellness Code format."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PURPLE_HEX = '7B52AB'
PURPLE_RGB = RGBColor(123, 82, 171)
ORANGE_RGB = RGBColor(240, 129, 30)
WHITE_RGB = RGBColor(255, 255, 255)
BLACK_RGB = RGBColor(30, 30, 30)
GRAY_RGB = RGBColor(100, 100, 100)
FONT = 'Calibri'

def shd(cell, hex_color):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear'); e.set(qn('w:color'), 'auto'); e.set(qn('w:fill'), hex_color)
    tcPr.append(e)

def get_or_add_tblPr(tbl_elem):
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr'); tbl_elem.insert(0, tblPr)
    return tblPr

def no_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl); tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}'); b.set(qn('w:val'), 'none'); tb.append(b)
    tblPr.append(tb)

def thin_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl); tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}'); b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D4C8F0'); tb.append(b)
    tblPr.append(tb)

def cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for name, val in [('top',top),('bottom',bottom),('left',left),('right',right)]:
        e = OxmlElement(f'w:{name}'); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); mar.append(e)
    tcPr.append(mar)

def spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr(); sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before)); sp.set(qn('w:after'), str(after)); pPr.append(sp)

def run(para, text, bold=False, italic=False, color=None, size=10.5):
    r = para.add_run(text); r.bold = bold; r.italic = italic
    r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color or BLACK_RGB
    return r

def section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1); no_borders(tbl)
    cell = tbl.cell(0, 0); shd(cell, PURPLE_HEX); cell_margins(cell, top=100, bottom=100, left=160, right=160)
    p = cell.paragraphs[0]; spacing(p); run(p, text, bold=True, color=WHITE_RGB, size=11.5)
    sp = doc.add_paragraph(); spacing(sp, before=0, after=40)

def two_col_table(doc, h1, h2, rows, c1=5.5, c2=11.0, warn=None):
    tbl = doc.add_table(rows=1+len(rows), cols=2); thin_borders(tbl)
    tbl.columns[0].width = Cm(c1); tbl.columns[1].width = Cm(c2)
    for ci, hdr in [(0,h1),(1,h2)]:
        cell = tbl.cell(0,ci); shd(cell, PURPLE_HEX); cell_margins(cell)
        p = cell.paragraphs[0]; spacing(p); run(p, hdr, bold=True, color=WHITE_RGB)
    warn = warn or []
    for i, (label, value) in enumerate(rows):
        c0 = tbl.cell(i+1,0); c1_ = tbl.cell(i+1,1)
        bg = 'EEE8F8' if i%2==0 else 'FFFFFF'
        shd(c0, bg); shd(c1_, bg); cell_margins(c0); cell_margins(c1_)
        p0 = c0.paragraphs[0]; spacing(p0); run(p0, label, bold=True, color=PURPLE_RGB)
        p1 = c1_.paragraphs[0]; spacing(p1)
        run(p1, value, bold=(i in warn), color=ORANGE_RGB if i in warn else BLACK_RGB)
    sp = doc.add_paragraph(); spacing(sp, before=0, after=80)


doc = Document()
for sec in doc.sections:
    sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5)
    sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)

p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; spacing(p,after=60)
run(p,'ВИТЯГ З АНКЕТИ ПАЦІЄНТА',bold=True,color=PURPLE_RGB,size=16)
p2 = doc.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; spacing(p2,after=40)
run(p2,'Скарги та порушення здоров\'я',italic=True,color=GRAY_RGB,size=11)
p3 = doc.add_paragraph(); p3.alignment=WD_ALIGN_PARAGRAPH.CENTER; spacing(p3,after=160)
run(p3,'Traets Magdalena',bold=True,color=ORANGE_RGB,size=13)

# 1. КОНТАКТНІ ДАНІ
section_header(doc,'1. КОНТАКТНІ ДАНІ')
two_col_table(doc,'Поле','Дані',[
    ('ПІБ','Traets Magdalena'),
    ('Дата народження','12.05.1962 (64 роки)'),
    ('Телефон','+45 24979495'),
    ('Email','marleen.traets@yahoo.com'),
    ('Адреса','Emdrupvej 24'),
    ('Рід занять','Господиня гест-хаусу'),
    ('Подорожі','Нігерія (2019, 2021), Угорщина (2 роки тому), Дубай (2 роки тому), США (давно)'),
])

# 2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ
section_header(doc,'2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ')
two_col_table(doc,'Показник','Значення',[
    ('Зріст','158 см'),
    ('Вага','55 кг'),
    ('Об\'єм талії','Не виміряно'),
    ('Об\'єм живота','Не виміряно'),
    ('Об\'єм грудей','Не виміряно'),
    ('Об\'єм стегон','Не виміряно'),
])

# 3. ОСНОВНІ СКАРГИ
section_header(doc,'3. ОСНОВНІ СКАРГИ')
two_col_table(doc,'Система / Орган','Скарга / Порушення',[
    ('Кровообіг / Кінцівки',
     'Руки та ноги завжди холодні (навіть влітку) ⚠️. Оніміння рук і ніг вночі та вранці — доводиться сідати прямо для відновлення кровообігу ⚠️. Поколювання, скутість. Дивні відчуття по тілу при початку руху після нерухомості'),
    ('Хребет / Спина',
     'Спина часто «виходить» — регулярно відвідує хіропрактика та фізіотерапевта. Куприк — падіння рік тому, запалення та зміщення, болі досі ⚠️. Часті судоми ніг вночі'),
    ('Суглоби / Кінцівки',
     'Болі в обох руках. Скутий палець правої стопи. Гомілковостоп — травма 3 роки тому, сильний набряк, не може бігати досі ⚠️. Рука — падіння 2007 р., досі відчуває скутість і хрускіт'),
    ('Панічні атаки / Нервова система',
     'Напади: тіло починає трястись як від холоду, поверхневе дихання, відчуття нестачі повітря ⚠️. Руки завжди холодні, синдром холодних кінцівок'),
    ('Гінекологія',
     'ВПЛ (високоризикований тип) — 2 біопсії (червень 2024 та січень 2025), зміни клітин не виявлено. Повторна перевірка через 1 рік ⚠️. Кіста (жіноча зона) — виявлена рік тому, при повторній біопсії відсутня. Менопауза — цикл відсутній'),
    ('Сон',
     'Засинає пізно (іноді після 00:00), важко засинає. Прокидається щонайменше 4 рази за ніч ⚠️. Прокидається не відпочившою'),
    ('ШКТ',
     'Здуття та гази (покращення при відмові від хліба). Іноді печія. Іноді відрижка. Стілець 1–2 рази/день, змішаний, іноді рідкий з частинками неперетравленої їжі'),
    ('Емоційний стан',
     'Стрес через особисту ситуацію в шлюбі. Пригніченість через фізичний стан. Перепади настрою. Тривога — інколи з причини'),
    ('COVID-19',
     'Двічі: 2020 та 2021 рр. Симптоми: кашель, жар, пітливість'),
    ('Менопауза',
     'Цикл відсутній. Сухість (менопауза). Вік менопаузи не уточнено'),
],warn=[0,1,2,3,4,5])

# 4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ
section_header(doc,'4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ')
two_col_table(doc,'Родич','Захворювання',[
    ('Батько','Проблеми зі шлунком, головні болі, носив окуляри'),
    ('Мати','Шлункова кровотеча (~40–50 років)'),
    ('Сестра 1','Захворювання крові — з 2013 р. не може повністю самостійно функціонувати ⚠️'),
    ('Сестра 2','Проблеми зі шкірою стоп та рук'),
],c1=6.0,c2=10.5,warn=[2])

# 5. ТРАВМИ ТА ОПЕРАТИВНІ ВТРУЧАННЯ
section_header(doc,'5. ТРАВМИ ТА ОПЕРАТИВНІ ВТРУЧАННЯ')
two_col_table(doc,'Рік / Вік','Втручання / Травма',[
    ('Дитинство (вік 5)','Травма голови — розтин шкіри, зашито лікарем'),
    ('2007','Падіння на руку — сильний набряк і синяк. Досі скутість і хрускіт'),
    ('3 роки тому','Травма гомілковостопу — набряк, синяк. Біг досі викликає біль ⚠️'),
    ('2024 (червень)','Біопсія (жіноча зона) — ВПЛ, зміни клітин не виявлено'),
    ('2025 (січень)','Біопсія повторна — норма. Спостереження через 1 рік ⚠️'),
    ('2025 (рік тому)','Падіння на куприк — запалення, зміщення. Болі досі ⚠️'),
    ('В анамнезі','Тендиніт. Видалено 4–5 зубів (давно) + 1 рік тому'),
],c1=4.5,c2=12.0,warn=[2,4,5])

# 6. ПОТОЧНА ТЕРАПІЯ / БАДи
section_header(doc,'6. ПОТОЧНА ТЕРАПІЯ / БАДи')
two_col_table(doc,'Препарат / БАД','Призначення',[
    ('Вітамін D3/K2','Кістки, імунітет'),
    ('Магній гліцинат','Нервова система, судоми'),
    ('Омега-3 (риб\'ячий жир)','Серцево-судинна система, запалення'),
    ('Омега-7','Слизові оболонки, суглоби'),
    ('Вітамін С','Імунітет'),
    ('Вітамін Е','Антиоксидант'),
    ('Вітамін А','Зір, імунітет'),
    ('Підтримка нервів (Dr. Berg)','Нейропатія, кровообіг'),
    ('Куркума (Turmeric)','Протизапальний ефект'),
    ('Мультивітаміни','Загальна підтримка'),
    ('Електроліти','Баланс мінералів'),
])

# 7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ
section_header(doc,'7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ')
two_col_table(doc,'Показник','Дані',[
    ('Кількість прийомів їжі','4 рази на день'),
    ('Сніданок (10:00–11:00)','Яйця з салатом/сиром/темним хлібом; смузі; тепла вівсянка з ягодами або бананом'),
    ('Обід (нерегулярно)','Салат, курячий бургер, сендвіч з сиром або шинкою, яйця'),
    ('Вечеря (19:00–20:00)','Картопля + овочі + м\'ясо; рис + овочі + риба; суп; іноді паста'),
    ('Перекуси','Горіхи, смузі, фрукти, йогурт з мюслі'),
    ('Порція','250–300 г'),
    ('Більше їсть','Ввечері ⚠️'),
    ('Вода на день','1000–1500 мл'),
    ('Кава','Не п\'є'),
    ('Спорт','Розтяжки та ходьба — щодня (намагається)'),
    ('Виключила','Білий цукор, молоко, білий хліб, каву'),
    ('Солодке','6/10 — знову з\'явилась тяга ⚠️'),
    ('Борошняне','6/10'),
    ('Алкоголь','Майже ніколи'),
],warn=[6,11])

# Footer
sp = doc.add_paragraph(); spacing(sp,before=200,after=0)
tbl = doc.add_table(rows=1,cols=2); no_borders(tbl)
tbl.columns[0].width=Cm(8); tbl.columns[1].width=Cm(8.5)
cell_margins(tbl.cell(0,0)); cell_margins(tbl.cell(0,1))
p0=tbl.cell(0,0).paragraphs[0]; run(p0,'Дата складання витягу: ________________',color=GRAY_RGB,size=9.5)
p1=tbl.cell(0,1).paragraphs[0]; p1.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run(p1,'Лікар: ________________',color=GRAY_RGB,size=9.5)

OUT='/home/user/Soloviova/Traets_Magdalena_Витяг.docx'
doc.save(OUT)
print(f'Saved: {OUT}')
