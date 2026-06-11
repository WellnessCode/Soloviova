"""Generate Uhrmann Josef patient extract in Smolych/Wellness Code format."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PURPLE_HEX = '7B52AB'
LAVENDER_HEX = 'EEE8F8'
PURPLE_RGB = RGBColor(123, 82, 171)
ORANGE_RGB = RGBColor(240, 129, 30)
WHITE_RGB = RGBColor(255, 255, 255)
BLACK_RGB = RGBColor(30, 30, 30)
GRAY_RGB = RGBColor(100, 100, 100)
FONT = 'Calibri'


def shd(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear')
    e.set(qn('w:color'), 'auto')
    e.set(qn('w:fill'), hex_color)
    tcPr.append(e)


def get_or_add_tblPr(tbl_elem):
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl_elem.insert(0, tblPr)
    return tblPr


def no_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl)
    tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'), 'none')
        tb.append(b)
    tblPr.append(tb)


def thin_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl)
    tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'D4C8F0')
        tb.append(b)
    tblPr.append(tb)


def cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in [('top',top),('bottom',bottom),('left',left),('right',right)]:
        e = OxmlElement(f'w:{name}')
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcPr.append(mar)


def spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'), str(after))
    pPr.append(sp)


def run(para, text, bold=False, italic=False, color=None, size=10.5):
    r = para.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color or BLACK_RGB
    return r


def section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    no_borders(tbl)
    cell = tbl.cell(0, 0)
    shd(cell, PURPLE_HEX)
    cell_margins(cell, top=100, bottom=100, left=160, right=160)
    p = cell.paragraphs[0]
    spacing(p)
    run(p, text, bold=True, color=WHITE_RGB, size=11.5)
    sp = doc.add_paragraph()
    spacing(sp, before=0, after=40)


def two_col_table(doc, col1_header, col2_header, rows_data, col1_w=5.5, col2_w=11.0, warn_rows=None):
    tbl = doc.add_table(rows=1+len(rows_data), cols=2)
    thin_borders(tbl)
    tbl.columns[0].width = Cm(col1_w)
    tbl.columns[1].width = Cm(col2_w)

    h0, h1 = tbl.cell(0,0), tbl.cell(0,1)
    shd(h0, PURPLE_HEX); shd(h1, PURPLE_HEX)
    cell_margins(h0); cell_margins(h1)
    p = h0.paragraphs[0]; spacing(p)
    run(p, col1_header, bold=True, color=WHITE_RGB)
    p = h1.paragraphs[0]; spacing(p)
    run(p, col2_header, bold=True, color=WHITE_RGB)

    warn_rows = warn_rows or []
    for i, (label, value) in enumerate(rows_data):
        c0 = tbl.cell(i+1, 0); c1 = tbl.cell(i+1, 1)
        bg = 'EEE8F8' if i % 2 == 0 else 'FFFFFF'
        shd(c0, bg); shd(c1, bg)
        cell_margins(c0); cell_margins(c1)
        p0 = c0.paragraphs[0]; spacing(p0)
        run(p0, label, bold=True, color=PURPLE_RGB)
        p1 = c1.paragraphs[0]; spacing(p1)
        val_color = ORANGE_RGB if i in warn_rows else BLACK_RGB
        val_bold = i in warn_rows
        run(p1, value, bold=val_bold, color=val_color)

    sp = doc.add_paragraph()
    spacing(sp, before=0, after=80)


# ── BUILD ─────────────────────────────────────────────────────────────────────
doc = Document()
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p, before=0, after=60)
run(p, 'ВИТЯГ З АНКЕТИ ПАЦІЄНТА', bold=True, color=PURPLE_RGB, size=16)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p2, before=0, after=40)
run(p2, 'Скарги та порушення здоров\'я', italic=True, color=GRAY_RGB, size=11)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p3, before=0, after=160)
run(p3, 'Uhrmann Josef', bold=True, color=ORANGE_RGB, size=13)

# 1. КОНТАКТНІ ДАНІ
section_header(doc, '1. КОНТАКТНІ ДАНІ')
two_col_table(doc, 'Поле', 'Дані', [
    ('ПІБ', 'Uhrmann Josef'),
    ('Дата народження', '12.03.1956 (69 років)'),
    ('Телефон', '01795299342'),
    ('Email', 'Josef.uhrmann@arcor.de'),
    ('Місце проживання', 'Мюнхен, Німеччина'),
    ('Рід занять', 'Електрик'),
    ('Подорожі', 'Данія (кілька разів), США (кілька разів), Аруба (2021)'),
])

# 2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ
section_header(doc, '2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ')
two_col_table(doc, 'Показник', 'Значення', [
    ('Зріст', '172 см'),
    ('Вага', '~89 кг  ⚠️ (бажане зниження ~10 кг)'),
    ('Об\'єм талії', 'Не виміряно'),
    ('Об\'єм грудей', 'Не виміряно'),
    ('Об\'єм стегон', 'Не виміряно'),
], warn_rows=[1])

# 3. ОСНОВНІ СКАРГИ
section_header(doc, '3. ОСНОВНІ СКАРГИ')
two_col_table(doc, 'Система / Орган', 'Скарга / Порушення', [
    ('Сон', 'Порушення сну — прокидається майже щоночі ⚠️. Засинає нестабільно. Прокидається невідпочившим. Часті сновидіння'),
    ('Метаболізм', 'Порушення метаболізму, надлишкова вага (~10 кг). Приймає Метформін (цукровий діабет 2 типу) ⚠️'),
    ('Артеріальний тиск', 'АТ 139/89 ⚠️. Приймає Севікар + Моксонідін'),
    ('Щитоподібна залоза', 'Гіперфункція (гіпертиреоз) ⚠️. Відчуття грудки в горлі'),
    ('Серцево-судинна', 'Задишка при фізичному навантаженні (інколи). Холестерин — приймає Розувастатин'),
    ('Травна система', 'Печія (інколи). Відрижка щодня. Здуття після їжі. Стілець 1–2 рази на день, залежить від їжі'),
    ('Нирки / Сечовипускання', 'Часте сечовипускання малими порціями ⚠️'),
    ('Суглоби / Стопа', 'Артрит лівого великого пальця ноги. Набряки стоп при фізичних навантаженнях'),
    ('Емоційний стан', 'Роздратування — майже щодня. Тривога — 3–4 рази на тиждень ⚠️'),
    ('COVID-19', 'Перехворів у 2020 р.'),
], warn_rows=[0, 1, 2, 3, 6, 8])

# 4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ
section_header(doc, '4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ')
two_col_table(doc, 'Родич', 'Захворювання', [
    ('—', 'Дані не надано'),
], col1_w=6.0, col2_w=10.5)

# 5. ТРАВМИ ТА ОПЕРАТИВНІ ВТРУЧАННЯ
section_header(doc, '5. ТРАВМИ ТА ОПЕРАТИВНІ ВТРУЧАННЯ')
two_col_table(doc, 'Рік / Вік', 'Втручання / Травма', [
    ('Вік 17', 'Травма правого коліна'),
    ('Вік 25', 'Травма підборіддя'),
    ('Вік 40', 'Травма носа'),
    ('1996', 'Запалення гомілки правої ноги — операція (загальна анестезія)'),
    ('Дата н/в', 'Операція на верхній щелепі'),
    ('Вік 62', 'Травма чола'),
    ('1996, 2025', 'Видалення постійних зубів'),
], col1_w=4.0, col2_w=12.5)

# 6. ПОТОЧНА ТЕРАПІЯ / БАДи
section_header(doc, '6. ПОТОЧНА ТЕРАПІЯ / БАДи')
two_col_table(doc, 'Препарат', 'Призначення', [
    ('Метформін', 'Цукровий діабет 2 типу ⚠️'),
    ('Севікар', 'Артеріальна гіпертензія'),
    ('Моксонідін', 'Артеріальна гіпертензія'),
    ('Розувастатин', 'Підвищений холестерин'),
], warn_rows=[0])

# 7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ
section_header(doc, '7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ')
two_col_table(doc, 'Показник', 'Дані', [
    ('Кількість прийомів їжі', '2–3 рази на день'),
    ('Порція', '300–400 г'),
    ('Основні продукти', 'М\'ясо, риба, паста'),
    ('Не їсть', 'Оливки, йогурт, артишоки'),
    ('Вода на день', '~1 літр ⚠️ (недостатньо)'),
    ('Кава', 'Натуральна з молоком, без цукру'),
    ('Спорт', 'Велосипед ~30 хв/день + вправи'),
], warn_rows=[4])

# Footer
sp = doc.add_paragraph()
spacing(sp, before=200, after=0)
tbl = doc.add_table(rows=1, cols=2)
no_borders(tbl)
tbl.columns[0].width = Cm(8)
tbl.columns[1].width = Cm(8.5)
cell_margins(tbl.cell(0,0)); cell_margins(tbl.cell(0,1))
p0 = tbl.cell(0,0).paragraphs[0]
run(p0, 'Дата складання витягу: ________________', color=GRAY_RGB, size=9.5)
p1 = tbl.cell(0,1).paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p1, 'Лікар: ________________', color=GRAY_RGB, size=9.5)

OUT = '/home/user/Soloviova/Uhrmann_Josef_Витяг.docx'
doc.save(OUT)
print(f'Saved: {OUT}')
