"""Generate Havryluk Oksana patient extract in Smolych format."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PURPLE_HEX = '7B52AB'
LAVENDER_HEX = 'EEE8F8'
PURPLE_RGB = RGBColor(123, 82, 171)
ORANGE_RGB = RGBColor(240, 129, 30)
WHITE_RGB = RGBColor(255, 255, 255)
BLACK_RGB = RGBColor(30, 30, 30)
GRAY_RGB = RGBColor(100, 100, 100)
FONT = 'Calibri'


def shd(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear')
    e.set(qn('w:color'), 'auto')
    e.set(qn('w:fill'), hex_color)
    tcPr.append(e)


def get_or_add_tblPr(tbl_elem):
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl_elem.insert(0, tblPr)
    return tblPr


def no_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl)
    tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'), 'none')
        tb.append(b)
    tblPr.append(tb)


def thin_borders(tbl):
    tblPr = get_or_add_tblPr(tbl._tbl)
    tb = OxmlElement('w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{s}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), 'D4C8F0')
        tb.append(b)
    tblPr.append(tb)


def cell_margins(cell, top=60, bottom=60, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for name, val in [('top',top),('bottom',bottom),('left',left),('right',right)]:
        e = OxmlElement(f'w:{name}')
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tcPr.append(mar)


def spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:before'), str(before))
    sp.set(qn('w:after'), str(after))
    pPr.append(sp)


def run(para, text, bold=False, italic=False, color=None, size=10.5):
    r = para.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color or BLACK_RGB
    return r


def section_header(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    no_borders(tbl)
    cell = tbl.cell(0, 0)
    shd(cell, PURPLE_HEX)
    cell_margins(cell, top=100, bottom=100, left=160, right=160)
    p = cell.paragraphs[0]
    spacing(p)
    run(p, text, bold=True, color=WHITE_RGB, size=11.5)
    sp = doc.add_paragraph()
    spacing(sp, before=0, after=40)


def two_col_table(doc, col1_header, col2_header, rows_data, col1_w=5.5, col2_w=11.0, warn_rows=None):
    tbl = doc.add_table(rows=1+len(rows_data), cols=2)
    thin_borders(tbl)
    tbl.columns[0].width = Cm(col1_w)
    tbl.columns[1].width = Cm(col2_w)

    # Header row
    h0, h1 = tbl.cell(0,0), tbl.cell(0,1)
    shd(h0, PURPLE_HEX); shd(h1, PURPLE_HEX)
    cell_margins(h0); cell_margins(h1)
    p = h0.paragraphs[0]; spacing(p)
    run(p, col1_header, bold=True, color=WHITE_RGB)
    p = h1.paragraphs[0]; spacing(p)
    run(p, col2_header, bold=True, color=WHITE_RGB)

    warn_rows = warn_rows or []
    for i, (label, value) in enumerate(rows_data):
        c0 = tbl.cell(i+1, 0); c1 = tbl.cell(i+1, 1)
        bg = 'EEE8F8' if i % 2 == 0 else 'FFFFFF'
        shd(c0, bg); shd(c1, bg)
        cell_margins(c0); cell_margins(c1)
        p0 = c0.paragraphs[0]; spacing(p0)
        run(p0, label, bold=True, color=PURPLE_RGB)
        p1 = c1.paragraphs[0]; spacing(p1)
        val_color = ORANGE_RGB if i in warn_rows else BLACK_RGB
        val_bold = i in warn_rows
        run(p1, value, bold=val_bold, color=val_color)

    sp = doc.add_paragraph()
    spacing(sp, before=0, after=80)


# ── BUILD ─────────────────────────────────────────────────────────────────────
doc = Document()
for section in doc.sections:
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p, before=0, after=60)
run(p, 'ВИТЯГ З АНКЕТИ ПАЦІЄНТА', bold=True, color=PURPLE_RGB, size=16)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p2, before=0, after=40)
run(p2, 'Скарги та порушення здоров\'я', italic=True, color=GRAY_RGB, size=11)

p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
spacing(p3, before=0, after=160)
run(p3, 'Гаврилюк Оксана', bold=True, color=ORANGE_RGB, size=13)

# 1. КОНТАКТНІ ДАНІ
section_header(doc, '1. КОНТАКТНІ ДАНІ')
two_col_table(doc, 'Поле', 'Дані', [
    ('ПІБ', 'Гаврилюк Оксана'),
    ('Дата народження', '12.06.1980 (45 років)'),
    ('Телефон', '+4917684318'),
    ('Email', 'Iriwa120680@gmail.com'),
    ('Місце проживання', 'Німеччина'),
    ('Рід занять', 'Обслуговування'),
])

# 2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ
section_header(doc, '2. АНТРОПОМЕТРИЧНІ ПОКАЗНИКИ')
two_col_table(doc, 'Показник', 'Значення', [
    ('Зріст', '160 см'),
    ('Вага', '94 кг  ⚠️ (бажана вага: -25–30 кг, цільова ~64–69 кг)'),
    ('Об\'єм талії', '100 см'),
    ('Об\'єм живота', '112 см'),
    ('Об\'єм грудей', '110 см'),
    ('Об\'єм стегон', '125 см'),
], warn_rows=[1])

# 3. ОСНОВНІ СКАРГИ
section_header(doc, '3. ОСНОВНІ СКАРГИ')
two_col_table(doc, 'Система / Орган', 'Скарга / Порушення', [
    ('Зайва вага', 'Надлишок 25–30 кг. Пов\'язує із захворюванням. Самостійні спроби результату не дали'),
    ('Набряки', 'Набрякає все тіло ⚠️'),
    ('Серцево-судинна', 'Перебої серцебиття — інколи. Задишка при фізичному навантаженні. Синці без причини — буває. Судини на ногах'),
    ('Суглоби', 'Біль у коліні'),
    ('Травна система', 'Відрижка майже постійно. Здуття живота — постійно ⚠️. Печія — інколи після певних продуктів'),
    ('Статева система', 'Багатовузлова міома тіла матки. Кіста правої молочної залози. Ліпома лівої молочної залози. Двобічна мастодинія ⚠️'),
    ('Менструальний цикл', 'Регулярний, болісний. Постійно приймає знеболюючі'),
    ('Психоемоційний стан', 'Перепади настрою протягом дня без причини. Роздратування — якщо є привід'),
    ('Сон', 'Прокидається невідпочившою. Важко прокидатись. Підйом о 5:30, відбій о 22:00'),
    ('Щитоподібна залоза', 'Кінцівки холодні — періодично'),
], warn_rows=[1, 5])

# 4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ
section_header(doc, '4. ОБТЯЖЕНИЙ СІМЕЙНИЙ АНАМНЕЗ')
two_col_table(doc, 'Родич', 'Захворювання', [
    ('Батько', 'Серцево-легенева недостатність'),
    ('Бабуся (по батьківській лінії)', 'Серцево-легенева недостатність'),
], col1_w=6.0, col2_w=10.5)

# 5. ОПЕРАТИВНІ ВТРУЧАННЯ В АНАМНЕЗІ
section_header(doc, '5. ОПЕРАТИВНІ ВТРУЧАННЯ В АНАМНЕЗІ')
two_col_table(doc, 'Рік', 'Операція', [
    ('1999', 'Мастит (загальна анестезія)'),
    ('2017', 'Видалення кісти (загальна анестезія)'),
], col1_w=3.0, col2_w=13.5)

# 6. ПОТОЧНА ТЕРАПІЯ / БАДи
section_header(doc, '6. ПОТОЧНА ТЕРАПІЯ / БАДи')
two_col_table(doc, 'Препарат', 'Призначення', [
    ('Петагма', 'Поточна терапія'),
    ('Іксіндол', 'Поточна терапія'),
    ('Магнос Леді', 'Поточна терапія'),
], col1_w=5.5, col2_w=11.0)

# 7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ
section_header(doc, '7. ХАРЧОВІ ЗВИЧКИ ТА ПОРУШЕННЯ')
two_col_table(doc, 'Показник', 'Дані', [
    ('Вода на день', '500–1000 мл ⚠️ (недостатньо)'),
    ('Кава', '3–4 чашки натуральної на день'),
    ('Кількість прийомів їжі', '3 рази + перекуси'),
    ('Сніданок (7:00–8:00)', 'Бутерброди, каші, комбіновані тарілки'),
    ('Обід', 'Салат з м\'ясом або бутерброди'),
    ('Вечеря (18:00–19:00)', 'Суп, салат, каша. Порція 250–300 г'),
    ('Перекуси', 'Горішки, банан, фрукти'),
    ('Солодке', '4/10 — помірно'),
    ('Борошняні вироби', '6/10 — підвищена тяга ⚠️'),
    ('Спорт', 'Відсутній ⚠️'),
    ('Алкоголь', 'Дуже рідко, вино'),
    ('Паління', 'Ні'),
], warn_rows=[0, 8, 9])

# Footer
sp = doc.add_paragraph()
spacing(sp, before=200, after=0)
tbl = doc.add_table(rows=1, cols=2)
no_borders(tbl)
tbl.columns[0].width = Cm(8)
tbl.columns[1].width = Cm(8.5)
cell_margins(tbl.cell(0,0)); cell_margins(tbl.cell(0,1))
p0 = tbl.cell(0,0).paragraphs[0]
run(p0, 'Дата складання витягу: ________________', color=GRAY_RGB, size=9.5)
p1 = tbl.cell(0,1).paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p1, 'Лікар: ________________', color=GRAY_RGB, size=9.5)

OUT = '/home/user/Soloviova/Гаврилюк_Оксана_Витяг.docx'
doc.save(OUT)
print(f'Saved: {OUT}')
